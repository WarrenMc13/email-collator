import csv
import io
from datetime import datetime
from typing import Sequence

COLUMNS = [
    ("date", "Date"),
    ("web_ref", "Web Ref"),
    ("enquirer_name", "Enquirer Name"),
    ("contact_number", "Contact Number"),
    ("enquirer_email", "Enquirer Email"),
    ("whatsapp_number", "WhatsApp Number"),
    ("message", "Message"),
    ("price", "Price"),
    ("property_type", "Property Type"),
    ("address", "Address"),
    ("property_url", "Property URL"),
]


def to_csv(enquiries: Sequence[dict], filepath: str) -> None:
    with open(filepath, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[k for k, _ in COLUMNS], extrasaction="ignore")
        writer.writerow({k: label for k, label in COLUMNS})
        writer.writerows(enquiries)


def to_excel(enquiries: Sequence[dict], filepath: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "Enquiries"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_idx, (key, label) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    alt_fill = PatternFill("solid", fgColor="D6E4F0")
    link_font = Font(color="1155CC", underline="single")

    for row_idx, enquiry in enumerate(enquiries, start=2):
        fill = alt_fill if row_idx % 2 == 0 else None
        for col_idx, (key, _) in enumerate(COLUMNS, start=1):
            value = enquiry.get(key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if fill:
                cell.fill = fill
            if key == "property_url" and value:
                cell.hyperlink = value
                cell.font = link_font
            if key == "whatsapp_number" and value:
                wa_url = f"https://wa.me/{value.replace(' ', '')}"
                cell.hyperlink = wa_url
                cell.value = value
                cell.font = link_font

    # Auto-size columns
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    ws.freeze_panes = "A2"
    wb.save(filepath)


def to_html(enquiries: Sequence[dict], filepath: str) -> None:
    rows_html = []
    for i, e in enumerate(enquiries):
        bg = "#f0f7ff" if i % 2 == 0 else "#ffffff"
        url = e.get("property_url", "")
        url_cell = f'<a href="{url}" target="_blank">{e.get("web_ref", "")}</a>' if url else e.get("web_ref", "")
        wa = e.get("whatsapp_number", "")
        wa_cell = (
            f'<a href="https://wa.me/{wa.replace(" ", "")}" target="_blank">{wa}</a>'
            if wa else ""
        )
        rows_html.append(f"""
        <tr style="background:{bg}">
            <td>{e.get('date','')}</td>
            <td>{url_cell}</td>
            <td>{e.get('enquirer_name','')}</td>
            <td>{e.get('contact_number','')}</td>
            <td><a href="mailto:{e.get('enquirer_email','')}">{e.get('enquirer_email','')}</a></td>
            <td>{wa_cell}</td>
            <td>{e.get('message','')}</td>
            <td>{e.get('price','')}</td>
            <td>{e.get('property_type','')}</td>
            <td>{e.get('address','')}</td>
        </tr>""")

    headers_html = "".join(f"<th>{label}</th>" for _, label in COLUMNS if _ != "property_url")

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Property24 Enquiries</title>
<style>
  body {{ font-family: Arial, sans-serif; padding: 20px; color: #222; }}
  h1 {{ color: #1F4E79; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
  th {{ background: #1F4E79; color: #fff; padding: 10px 8px; text-align: left; }}
  td {{ padding: 8px; border-bottom: 1px solid #dce6f0; vertical-align: top; }}
  a {{ color: #1155CC; }}
  .meta {{ color: #666; font-size: 12px; margin-bottom: 16px; }}
</style>
</head>
<body>
<h1>Property24 Enquiries</h1>
<p class="meta">Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} &nbsp;|&nbsp; Total: {len(enquiries)}</p>
<table>
  <thead><tr>{headers_html}</tr></thead>
  <tbody>{''.join(rows_html)}</tbody>
</table>
</body>
</html>"""

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(html)
